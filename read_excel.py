import json
try:
    from openpyxl import load_workbook
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl', '-q'])
    from openpyxl import load_workbook

wb = load_workbook('VisionOS_MVP.xlsx', data_only=True)
print("=== Sheet names ===")
print(wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n=== Sheet: {sheet_name} ===")
    print(f"Rows: {ws.max_row}, Cols: {ws.max_column}")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=False):
        row_data = []
        for cell in row:
            row_data.append(str(cell.value) if cell.value is not None else "")
        print(" | ".join(row_data))
